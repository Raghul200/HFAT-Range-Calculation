
import streamlit as st
import pandas as pd
from openpyxl import load_workbook
from io import BytesIO

st.set_page_config(page_title="Infinity IO Extractor", layout="wide")
st.title("Infinity IO Extractor")

uploaded_file = st.file_uploader("Upload Excel File", type=["xlsx","xlsm"])

if uploaded_file:
    wb = load_workbook(uploaded_file, data_only=True)
    ws = wb.active

    output = []
    headers = [ws["B1"].value, ws["C1"].value, ws["D1"].value, ws["F1"].value,
               ws["L1"].value, ws["M1"].value, ws["N1"].value, ws["O1"].value,
               "Range 1","Range 2","Range 3"]

    for row in ws.iter_rows(min_row=2):
        if row[3].value in ("InfinityInput", "InfinityOutput"):
            if any((row[c].value not in (None,"")) for c in [11,12,13,14]):
                rec = [row[1].value,row[2].value,row[3].value,row[5].value,
                       row[11].value,row[12].value,row[13].value,row[14].value]

                r1=r2=r3=""
                if row[3].value=="InfinityInput":
                    low=float(row[14].value or 0)
                    high=float(row[13].value or 0)
                    total=high-low
                    x=(total*0.58)/100

                    r1=f"{low-x:.2f} to {low+x:.2f}"
                    mid=low+(total*0.5)
                    r2=f"{mid-x:.2f} to {mid+x:.2f}"
                    r3=f"{high-x:.2f} to {high+x:.2f}"

                rec.extend([r1,r2,r3])
                output.append(rec)

    df=pd.DataFrame(output,columns=headers)

    bio=BytesIO()
    with pd.ExcelWriter(bio,engine="openpyxl") as writer:
        df.to_excel(writer,index=False,sheet_name="Infinity_IO")

    st.success(f"{len(df)} rows extracted")
    st.download_button("Download Output",
                       bio.getvalue(),
                       "Infinity_IO.xlsx",
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
